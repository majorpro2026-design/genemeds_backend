from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from etl.utils.validation import normalize_text, row_has_data


def discover_workbooks(root: Path | str = "cpic") -> list[Path]:
    base = Path(root)
    if not base.exists():
        return []
    return sorted(
        path
        for path in base.rglob("*.xlsx")
        if not path.name.startswith("~$")
    )


def infer_drug_name(path: Path) -> str | None:
    parts = path.parent.name.split("_")
    if not parts:
        return None
    return parts[0].strip().replace("-", " ").title()


def _sheet_rows(sheet, max_rows: int = 10, max_cols: int = 30) -> list[list[object]]:
    rows: list[list[object]] = []
    for row in sheet.iter_rows(min_row=1, max_row=max_rows, max_col=max_cols, values_only=True):
        values = list(row)
        if row_has_data(values):
            rows.append(values)
    return rows


def classify_workbook(path: Path) -> str | None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    filename = path.name.lower()

    if "allele_definition" in filename:
        return "allele_definition"
    if "diplotype" in filename:
        return "diplotype_phenotype"
    if "recommendation" in filename:
        return "recommendation"

    for sheet in workbook.worksheets:
        rows = _sheet_rows(sheet)
        headers = [normalize_text(value) or "" for value in rows[0]] if rows else []
        header_blob = " | ".join(headers).lower()
        if "therapeutic recommendation" in header_blob and "classification of recommendation" in header_blob:
            return "recommendation"
        if "coded diplotype/phenotype summary" in header_blob:
            return "diplotype_phenotype"
        if rows and isinstance(rows[0][0], str) and rows[0][0].lower().startswith("gene:"):
            return "allele_definition"
    return None


