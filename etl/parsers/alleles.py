from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from etl.parsers.common import iter_nonempty_rows
from etl.utils.validation import normalize_text


@dataclass(frozen=True)
class AlleleRow:
    star_allele: str
    functionality: str | None = None
    activity_value: float | None = None


@dataclass(frozen=True)
class AlleleWorkbook:
    path: Path
    sheet_name: str
    gene_symbol: str
    rows: list[AlleleRow]


def parse_allele_definition_workbook(path: Path) -> AlleleWorkbook | None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        first_cell = normalize_text(sheet["A1"].value)
        if not first_cell or not first_cell.lower().startswith("gene:"):
            continue

        gene_symbol = first_cell.split(":", 1)[1].strip()
        rows: list[AlleleRow] = []
        allele_start_row = None
        for row_number, values in iter_nonempty_rows(sheet, start_row=1):
            header = normalize_text(values[0])
            if header and header.endswith("Allele"):
                allele_start_row = row_number + 1
                break

        if allele_start_row is None:
            return None

        for _, values in iter_nonempty_rows(sheet, start_row=allele_start_row):
            star_allele = normalize_text(values[0])
            if not star_allele or not star_allele.startswith("*"):
                continue
            rows.append(AlleleRow(star_allele=star_allele))

        return AlleleWorkbook(path=path, sheet_name=sheet.title, gene_symbol=gene_symbol, rows=rows)
    return None
