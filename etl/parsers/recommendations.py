from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from etl.parsers.common import find_header_row, iter_nonempty_rows
from etl.utils.validation import normalize_text


@dataclass(frozen=True)
class RecommendationCondition:
    gene_symbol: str
    phenotype_name: str


@dataclass(frozen=True)
class RecommendationRow:
    conditions: list[RecommendationCondition]
    recommendation_text: str
    implication: str | None
    recommendation_strength: str | None
    classification: str | None
    source: str = "CPIC"
    cpic_guideline_version: str | None = None
    evidence_level: str | None = None
    recommendation_order: int = 1


@dataclass(frozen=True)
class RecommendationWorkbook:
    path: Path
    sheet_name: str
    rows: list[RecommendationRow]


def _extract_conditions(headers: list[str], values: list[object]) -> list[RecommendationCondition]:
    conditions: list[RecommendationCondition] = []
    for index, header in enumerate(headers):
        if "phenotype" not in header.lower():
            continue
        value = normalize_text(values[index]) if index < len(values) else None
        if not value:
            continue
        gene_symbol = header.split("Phenotype", 1)[0].strip()
        conditions.append(RecommendationCondition(gene_symbol=gene_symbol, phenotype_name=value))
    return conditions


def parse_recommendation_workbook(path: Path) -> RecommendationWorkbook | None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        header = find_header_row(
            sheet,
            required_terms=["therapeutic recommendation", "classification of recommendation"],
            scan_rows=10,
        )
        if header is None:
            continue

        rows: list[RecommendationRow] = []
        for row_number, values in iter_nonempty_rows(sheet, start_row=header.header_row + 1):
            conditions = _extract_conditions(header.headers, values)
            recommendation_text = None
            implication = None
            recommendation_strength = None
            classification = None

            for index, column in enumerate(header.headers):
                value = normalize_text(values[index]) if index < len(values) else None
                if not value:
                    continue
                lowered = column.lower()
                if lowered == "therapeutic recommendation":
                    recommendation_text = value
                elif "implication" in lowered:
                    implication = value
                elif "classification of recommendation" in lowered:
                    classification = value
                elif "recommendation strength" in lowered:
                    recommendation_strength = value
                elif "classification" == lowered:
                    classification = value

            if not conditions or recommendation_text is None:
                continue

            rows.append(
                RecommendationRow(
                    conditions=conditions,
                    recommendation_text=recommendation_text,
                    implication=implication,
                    recommendation_strength=recommendation_strength,
                    classification=classification,
                    recommendation_order=row_number,
                )
            )

        return RecommendationWorkbook(path=path, sheet_name=sheet.title, rows=rows)
    return None
