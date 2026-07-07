from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from etl.parsers.common import find_header_row, iter_nonempty_rows
from etl.utils.validation import normalize_numeric, normalize_text


@dataclass(frozen=True)
class DiplotypeRow:
    gene_symbol: str
    diplotype: str
    allele1: str
    allele2: str
    activity_score: float | None
    phenotype_name: str
    phenotype_description: str | None


@dataclass(frozen=True)
class DiplotypeWorkbook:
    path: Path
    sheet_name: str
    rows: list[DiplotypeRow]


def _split_diplotype(diplotype: str) -> tuple[str, str] | None:
    parts = [part.strip() for part in diplotype.split("/")]
    if len(parts) != 2:
        return None
    return parts[0], parts[1]


def parse_diplotype_workbook(path: Path) -> DiplotypeWorkbook | None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        header = find_header_row(
            sheet,
            required_terms=["diplotype", "phenotype", "summary"],
            scan_rows=10,
        )
        if header is None:
            continue

        rows: list[DiplotypeRow] = []
        gene_symbol = header.headers[0].split("Diplotype", 1)[0].strip()
        for _, values in iter_nonempty_rows(sheet, start_row=header.header_row + 1):
            diplotype = normalize_text(values[0])
            if not diplotype or not diplotype.startswith("*"):
                continue
            parsed = _split_diplotype(diplotype)
            if parsed is None:
                continue
            allele1, allele2 = parsed
            activity_score = normalize_numeric(values[1]) if len(values) > 1 else None
            phenotype_name = normalize_text(values[2]) if len(values) > 2 else None
            if phenotype_name is None:
                continue
            rows.append(
                DiplotypeRow(
                    gene_symbol=gene_symbol,
                    diplotype=diplotype,
                    allele1=allele1,
                    allele2=allele2,
                    activity_score=activity_score,
                    phenotype_name=phenotype_name,
                    phenotype_description=normalize_text(values[3]) if len(values) > 3 else None,
                )
            )

        return DiplotypeWorkbook(path=path, sheet_name=sheet.title, rows=rows)
    return None

